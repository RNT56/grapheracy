from __future__ import annotations

from io import BytesIO
from pathlib import PurePosixPath
from zipfile import BadZipFile, ZipFile

from docx import Document
from openpyxl import load_workbook
from pptx import Presentation

from graphview_api.ingestion import extract_pdf_text


ALLOWED_UPLOAD_SUFFIXES = {".txt", ".md", ".markdown", ".json", ".csv", ".pdf", ".docx", ".xlsx", ".pptx"}


def validate_uploaded_payload(filename: str, content_type: str, payload: bytes) -> None:
    suffix = PurePosixPath(filename).suffix.lower()
    if suffix not in ALLOWED_UPLOAD_SUFFIXES:
        raise ValueError(f"Unsupported upload type: {suffix or 'missing extension'}")
    if suffix == ".pdf" and not payload.startswith(b"%PDF-"):
        raise ValueError("PDF upload has an invalid file signature")
    if suffix in {".docx", ".xlsx", ".pptx"}:
        if not payload.startswith(b"PK"):
            raise ValueError("Office upload has an invalid file signature")
        try:
            with ZipFile(BytesIO(payload)) as archive:
                compressed = max(1, sum(item.compress_size for item in archive.infolist()))
                expanded = sum(item.file_size for item in archive.infolist())
                if expanded > 200_000_000 or expanded / compressed > 100:
                    raise ValueError("Office upload exceeds safe expansion limits")
        except BadZipFile as error:
            raise ValueError("Office upload is not a valid archive") from error
    elif suffix != ".pdf":
        if b"\x00" in payload[:8192]:
            raise ValueError("Text upload contains binary data")
        try:
            payload.decode("utf-8")
        except UnicodeDecodeError as error:
            raise ValueError("Text upload must be UTF-8") from error


def extract_uploaded_text(filename: str, content_type: str, payload: bytes) -> tuple[str, str]:
    validate_uploaded_payload(filename, content_type, payload)
    suffix = PurePosixPath(filename).suffix.lower()
    if content_type == "application/pdf" or suffix == ".pdf":
        return "pdf", extract_pdf_text(payload)
    if suffix == ".docx":
        document = Document(BytesIO(payload))
        paragraphs = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
        tables = [" | ".join(cell.text for cell in row.cells) for table in document.tables for row in table.rows]
        return "ops-document", "\n".join([*paragraphs, *tables])
    if suffix == ".xlsx":
        workbook = load_workbook(BytesIO(payload), read_only=True, data_only=True)
        lines = []
        for sheet in workbook.worksheets:
            lines.append(f"# {sheet.title}")
            lines.extend(" | ".join("" if value is None else str(value) for value in row) for row in sheet.iter_rows(values_only=True))
        return "ops-document", "\n".join(lines)
    if suffix == ".pptx":
        presentation = Presentation(BytesIO(payload))
        lines = [shape.text for slide in presentation.slides for shape in slide.shapes if hasattr(shape, "text") and shape.text.strip()]
        return "ops-document", "\n".join(lines)
    return ("markdown" if suffix in {".md", ".markdown"} else "text"), payload.decode("utf-8", errors="replace")
