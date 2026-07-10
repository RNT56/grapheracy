from io import BytesIO

from docx import Document
from openpyxl import Workbook
from pptx import Presentation

import pytest

from graphview_api.upload_extraction import extract_uploaded_text, validate_uploaded_payload


def test_office_upload_extractors_cover_word_excel_and_powerpoint() -> None:
    word_stream = BytesIO()
    word = Document()
    word.add_paragraph("Word evidence")
    word.save(word_stream)

    excel_stream = BytesIO()
    workbook = Workbook()
    workbook.active.title = "Signals"
    workbook.active.append(["Signal", "Severity"])
    workbook.active.append(["Latency", "High"])
    workbook.save(excel_stream)

    presentation_stream = BytesIO()
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "Review summary"
    presentation.save(presentation_stream)

    assert "Word evidence" in extract_uploaded_text("evidence.docx", "application/octet-stream", word_stream.getvalue())[1]
    assert "Latency | High" in extract_uploaded_text("signals.xlsx", "application/octet-stream", excel_stream.getvalue())[1]
    assert "Review summary" in extract_uploaded_text("review.pptx", "application/octet-stream", presentation_stream.getvalue())[1]


def test_upload_validation_rejects_binary_disguises_and_unsupported_extensions() -> None:
    with pytest.raises(ValueError, match="invalid file signature"):
        validate_uploaded_payload("malicious.pdf", "application/pdf", b"not a pdf")
    with pytest.raises(ValueError, match="binary data"):
        validate_uploaded_payload("malicious.txt", "text/plain", b"text\x00binary")
    with pytest.raises(ValueError, match="Unsupported upload type"):
        validate_uploaded_payload("payload.exe", "application/octet-stream", b"MZ")
